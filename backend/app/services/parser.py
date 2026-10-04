from io import BytesIO
from pypdf import PdfReader
from docx import Document


def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))
    text_parts = []
    for page in reader.pages:
        text_parts.append(page.extract_text() or "")
    return "\n".join(text_parts)


def _extract_docx(content: bytes) -> str:
    document = Document(BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]

    # Also read text from tables because resumes often use table layouts.
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" ".join(cell.text for cell in row.cells))

    return "\n".join(paragraphs)


def extract_resume_text(content: bytes, extension: str) -> str:
    if extension == ".pdf":
        return _extract_pdf(content)
    if extension == ".docx":
        return _extract_docx(content)
    raise ValueError("Unsupported file type.")
